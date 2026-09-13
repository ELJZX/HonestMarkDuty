"""Формирование Excel-файла чеклиста оборудования «Честный знак»."""
from __future__ import annotations

from io import BytesIO

from django.core.files.base import ContentFile
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from checklists.models import ChecklistCheck, ChecklistGroup

TITLE = 'Чек лист технического осмотра оборудования "Честный знак"'
SCORE_FILL = {1: "FF0000", 2: "FFFF00", 3: "00B050"}
SCORE_FONT = {1: "FFFFFF", 2: "000000", 3: "FFFFFF"}
LEGEND = [
    (1, "Неудовлетворительно, критические несоответствия"),
    (2, "Удовлетворительно, присутствуют недочеты"),
    (3, "Оборудование настроено, обслужено и готово к дальнейшей работе"),
]


def _machines_layout():
    groups = list(ChecklistGroup.objects.prefetch_related("machines").order_by("sort_order", "name"))
    machines = []
    spans = []
    column = 3
    for group in groups:
        group_machines = list(group.machines.order_by("sort_order", "name"))
        if not group_machines:
            continue
        start = column
        for machine in group_machines:
            machines.append(machine)
            column += 1
        spans.append((group.name, start, column - 1))
    return machines, spans, max(column - 1, 2)


def build_checklist_workbook(checklist) -> Workbook:
    machines, spans, last_col = _machines_layout()
    checks = list(ChecklistCheck.objects.order_by("sort_order", "id"))

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Чеклист"

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    title_cell = sheet.cell(1, 1, TITLE)
    title_cell.font = Font(size=15, bold=True)
    title_cell.alignment = center
    sheet.row_dimensions[1].height = 26

    sheet.merge_cells("A2:B2")
    date_cell = sheet.cell(2, 1, f"Дата: {checklist.date:%d.%m.%Y}")
    date_cell.font = Font(bold=True)
    date_cell.alignment = center
    if last_col >= 3:
        sheet.merge_cells(start_row=2, start_column=3, end_row=2, end_column=last_col)
        head_cell = sheet.cell(2, 3, "Цех / автомат")
        head_cell.font = Font(bold=True)
        head_cell.alignment = center

    sheet.merge_cells("A3:A4")
    sheet.cell(3, 1, "№").font = Font(bold=True)
    sheet.cell(3, 1).alignment = center
    sheet.merge_cells("B3:B4")
    sheet.cell(3, 2, "Проводимые работы и проверки").font = Font(bold=True)
    sheet.cell(3, 2).alignment = center

    for name, start, end in spans:
        if end > start:
            sheet.merge_cells(start_row=3, start_column=start, end_row=3, end_column=end)
        cell = sheet.cell(3, start, name)
        cell.font = Font(bold=True)
        cell.alignment = center

    for index, machine in enumerate(machines, start=3):
        cell = sheet.cell(4, index, machine.name)
        cell.font = Font(bold=True, size=9)
        cell.alignment = center

    for row in (2, 3, 4):
        for column in range(1, last_col + 1):
            sheet.cell(row, column).border = border

    sheet.column_dimensions["A"].width = 5
    sheet.column_dimensions["B"].width = 44
    for column in range(3, last_col + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 8.5

    mileage = {item.machine_id: item.value for item in checklist.mileages.all()}
    row = 5
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    mileage_label = sheet.cell(row, 1, "Пробег печатающей головки (км)")
    mileage_label.font = Font(bold=True, size=10)
    mileage_label.alignment = left
    for index, machine in enumerate(machines, start=3):
        value = mileage.get(machine.id)
        cell = sheet.cell(row, index, float(value) if value is not None else None)
        cell.alignment = center
    for column in range(1, last_col + 1):
        sheet.cell(row, column).border = border

    results = {
        (item.machine_id, item.check_item_id): item.score for item in checklist.results.all()
    }
    row = 6
    for number, check in enumerate(checks, start=1):
        sheet.cell(row, 1, number).alignment = center
        sheet.cell(row, 2, check.name).alignment = left
        for index, machine in enumerate(machines, start=3):
            score = results.get((machine.id, check.id))
            cell = sheet.cell(row, index, score if score else None)
            cell.alignment = center
            if score:
                cell.fill = PatternFill("solid", fgColor=SCORE_FILL.get(score, "FFFFFF"))
                cell.font = Font(bold=True, color=SCORE_FONT.get(score, "000000"))
        for column in range(1, last_col + 1):
            sheet.cell(row, column).border = border
        row += 1

    row += 1
    sheet.cell(row, 2, "Выполнил:  Специалист по цифровой маркировке").font = Font(bold=True)
    row += 2
    sheet.cell(row, 2, "Проверил:  Ведущий инженер по цифровой маркировке").font = Font(bold=True)

    legend_col = max(3, last_col - 1)
    legend_row = row - 3
    for score, text in LEGEND:
        cell = sheet.cell(legend_row, legend_col, score)
        cell.fill = PatternFill("solid", fgColor=SCORE_FILL[score])
        cell.font = Font(bold=True, color=SCORE_FONT[score])
        cell.alignment = center
        sheet.cell(legend_row, legend_col + 1, text).alignment = left
        legend_row += 1

    return workbook


def save_checklist_file(checklist):
    workbook = build_checklist_workbook(checklist)
    buffer = BytesIO()
    workbook.save(buffer)
    filename = f"checklist_{checklist.date:%Y%m%d}_{checklist.pk}.xlsx"
    checklist.file.save(filename, ContentFile(buffer.getvalue(), name=filename), save=False)
    checklist.save()
    return checklist.file

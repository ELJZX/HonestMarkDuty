"""Формирование Excel-файла чеклиста принтеров Markem Image 9450."""

from __future__ import annotations

from io import BytesIO

from django.core.files.base import ContentFile
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from checklists.models import MarkemParameter, MarkemPrinter

TITLE = "Чек лист технического осмотра и обслуживания принтеров Markem Image 9450"


def _layout():
    printers = list(MarkemPrinter.objects.order_by("sort_order", "id"))
    parameters = list(MarkemParameter.objects.order_by("sort_order", "id"))
    return printers, parameters, max(2 + len(printers), 3)


def build_markem_workbook(checklist) -> Workbook:
    printers, parameters, last_col = _layout()

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Markem 9450"

    # Заголовок
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    title_cell = sheet.cell(1, 1, TITLE)
    title_cell.font = Font(size=15, bold=True)
    title_cell.alignment = center
    sheet.row_dimensions[1].height = 26

    # Дата
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=2)
    date_cell = sheet.cell(2, 3, f"Дата: {checklist.date:%d.%m.%Y}")
    date_cell.font = Font(size=13, bold=True)
    date_cell.alignment = center
    if last_col >= 3:
        sheet.merge_cells(start_row=2, start_column=3, end_row=2, end_column=last_col)

    # Шапка таблицы
    sheet.merge_cells(start_row=3, start_column=1, end_row=4, end_column=1)
    sheet.cell(3, 1, "№").font = Font(bold=True)
    sheet.cell(3, 1).alignment = center
    sheet.merge_cells(start_row=3, start_column=2, end_row=4, end_column=2)
    sheet.cell(3, 2, "Проводимые работы и проверка параметров").font = Font(bold=True)
    sheet.cell(3, 2).alignment = center

    for index, printer in enumerate(printers, start=3):
        if index <= last_col:
            number_cell = sheet.cell(3, index, index - 2)
            number_cell.font = Font(bold=True)
            number_cell.alignment = center
        serial_cell = sheet.cell(4, index, printer.name)
        serial_cell.font = Font(bold=True, size=9)
        serial_cell.alignment = center

    for row in (2, 3, 4):
        for column in range(1, last_col + 1):
            sheet.cell(row, column).border = border

    sheet.column_dimensions["A"].width = 6
    sheet.column_dimensions["B"].width = 46
    for column in range(3, last_col + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 13

    values = {
        (item.printer_id, item.parameter_id): item.value for item in checklist.values.all()
    }

    row = 5
    for number, parameter in enumerate(parameters, start=1):
        sheet.cell(row, 1, number).alignment = center
        sheet.cell(row, 2, parameter.name).alignment = left
        for index, printer in enumerate(printers, start=3):
            cell = sheet.cell(row, index, values.get((printer.id, parameter.id)) or None)
            cell.alignment = center
        for column in range(1, last_col + 1):
            sheet.cell(row, column).border = border
        row += 1

    # Подписи
    row += 1
    sheet.cell(row, 2, "Выполнил:  Специалист по цифровой маркировке").font = Font(bold=True)
    row += 2
    sheet.cell(row, 2, "Проверил:  Ведущий инженер по цифровой маркировке").font = Font(bold=True)

    return workbook


def save_markem_file(checklist):
    workbook = build_markem_workbook(checklist)
    buffer = BytesIO()
    workbook.save(buffer)
    filename = f"markem_{checklist.date:%Y%m%d}_{checklist.pk}.xlsx"
    checklist.file.save(filename, ContentFile(buffer.getvalue(), name=filename), save=False)
    checklist.save()
    return checklist.file

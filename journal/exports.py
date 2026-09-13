"""Формирование единого Excel-архива сменного журнала (по утверждённой форме)."""
from __future__ import annotations

from collections import OrderedDict
from io import BytesIO

from django.core.files.base import ContentFile
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from journal.models import JournalEntry, JournalExport

TITLE = "Сменный журнал специалистов по цифровой маркировке"

HEADERS = [
    ("Дата", 14),
    ("Дежурный специалист", 26),
    ("Время", 12),
    ("Оборудование / Линия", 24),
    ("Время простоя", 16),
    ("Действие / Задача", 46),
    ("Решение (комментарий)", 46),
    ("Печ. головка", 14),
    ("Пробег (км)", 12),
]

THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="1F1F1F")
YELLOW_FILL = PatternFill("solid", fgColor="FFFF00")


def local_dt(value):
    return timezone.localtime(value) if value else None


def format_time(value) -> str:
    dt = local_dt(value)
    return dt.strftime("%H:%M") if dt else ""


def group_entries(entries) -> list[list[JournalEntry]]:
    """Группирует записи по смене (или по дате и специалисту) — для объединённых ячеек."""
    groups: "OrderedDict[tuple, list[JournalEntry]]" = OrderedDict()
    for entry in entries:
        if entry.shift_id:
            key = ("shift", entry.shift_id)
        else:
            dt = local_dt(entry.occurred_at)
            key = ("day", dt.date() if dt else None, entry.specialist_name)
        groups.setdefault(key, []).append(entry)
    return list(groups.values())


def build_workbook(entries) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Сменный журнал"
    columns = len(HEADERS)

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    title_cell = sheet.cell(row=1, column=1, value=TITLE)
    title_cell.font = Font(size=16, bold=True)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 26

    header_row = 2
    for index, (header, width) in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=header_row, column=index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[header_row].height = 34

    row = header_row + 1
    for group in group_entries(list(entries)):
        start_row = row
        for entry in group:
            values = [
                None,
                None,
                format_time(entry.occurred_at),
                entry.equipment_line,
                entry.downtime,
                entry.action_task,
                entry.solution,
                entry.print_head,
                float(entry.mileage) if entry.mileage is not None else None,
            ]
            for index, value in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=index, value=value)
                cell.border = BORDER
                cell.alignment = Alignment(
                    vertical="center" if index < 3 else "top",
                    wrap_text=index in (4, 5, 6, 7),
                )
            row += 1

        end_row = row - 1
        first = group[0]
        if end_row > start_row:
            sheet.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
            sheet.merge_cells(start_row=start_row, start_column=2, end_row=end_row, end_column=2)

        date_cell = sheet.cell(row=start_row, column=1, value=local_dt(first.occurred_at).date())
        date_cell.number_format = "DD.MM.YYYY"
        specialist_cell = sheet.cell(row=start_row, column=2, value=first.specialist_name)
        for cell in (date_cell, specialist_cell):
            cell.fill = YELLOW_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BORDER

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    return workbook


def _to_content(workbook: Workbook, filename: str) -> ContentFile:
    buffer = BytesIO()
    workbook.save(buffer)
    return ContentFile(buffer.getvalue(), name=filename)


def export_full_journal(user=None, shift=None, entries=None) -> JournalExport | None:
    """Формирует/обновляет единый Excel-архив журнала со всеми записями."""
    if entries is None:
        entries = JournalEntry.objects.select_related(
            "shift", "specialist", "shift__opened_by"
        ).order_by("occurred_at")
    entries = list(entries)
    if not entries:
        return None

    filename = "smennyy_zhurnal.xlsx"
    workbook = build_workbook(entries)
    content = _to_content(workbook, filename)

    export = JournalExport.objects.filter(is_full=True).first()
    if export is None:
        export = JournalExport(is_full=True)
    else:
        if export.file:
            export.file.delete(save=False)

    export.shift = shift or export.shift
    export.entries_count = len(entries)
    export.period_start = entries[0].occurred_at
    export.period_end = entries[-1].occurred_at
    if user is not None:
        export.created_by = user
    export.file.save(filename, content, save=False)
    export.save()
    return export

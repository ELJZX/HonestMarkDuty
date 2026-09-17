"""Формирование документов по цехам (техническое заключение, служебная записка).

Файл-образец ищется в папке ``documents/samples/`` по имени ``<код>_<тип>.docx``
(например, ``ceh1_tz.docx``). Если образца нет — формируется документ-заглушка.

В образцах поддерживаются плейсхолдеры ``{{ ... }}``, дата ``«ДАТА» «МЕСЯЦ» «ГОД»``
(→ ``«16» сентября 2026``) и данные сменного специалиста: ``«ФИО»``,
``«ФАМИЛИЯ/ИНИЦИАЛЫ»``, строка ``От: ...``, должность в строке ``Должность:``
после ``От:`` и подпись вида ``<должность>  _______ <ФИО>``.
"""
from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.utils import timezone
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from documents.services import MONTHS_RU, date_long, render_text

DATE_TOKEN_RE = re.compile(r"«\s*(ДАТА|МЕСЯЦ|ГОД)\s*»")
FROM_LINE_RE = re.compile(r"(От:\s*).*$")
POSITION_RE = re.compile(r"^(\s*Должность:\s*).*$")
SIGNATURE_LINE_RE = re.compile(r"^(?P<before>.*?)(?P<us>_{3,})(?P<after>.*)$")

WORKSHOP_DOCUMENTS = [
    {"name": "Цех №1", "code": "ceh1"},
    {"name": "Кисломолочный цех", "code": "kmc"},
    {"name": "Цельномолочный цех", "code": "cm"},
    {"name": "Цех стерильного молока", "code": "csm"},
    {"name": "Творожный цех", "code": "tv"},
    {"name": "Малыш моцарелла", "code": "mc"},
    {"name": "Малыш сырки", "code": "ms"},
]

DOC_KINDS = {
    "tz": "Техническое заключение",
    "sl": "Служебная записка",
}

WORKSHOP_BY_CODE = {item["code"]: item["name"] for item in WORKSHOP_DOCUMENTS}


def samples_dir() -> Path:
    configured = getattr(settings, "DOCUMENT_SAMPLES_DIR", None)
    if configured:
        return Path(configured)
    return Path(__file__).resolve().parent / "samples"


def sample_path(code: str, kind: str) -> Path:
    return samples_dir() / f"{code}_{kind}.docx"


def short_name(user) -> str:
    """Фамилия и инициалы: «Бобров М.А.»."""
    if user is None:
        return ""
    last = (getattr(user, "last_name", "") or "").strip()
    first = (getattr(user, "first_name", "") or "").strip()
    patronymic = (getattr(user, "patronymic", "") or "").strip()
    if last:
        initials = "".join(f"{part[0].upper()}." for part in (first, patronymic) if part)
        return f"{last} {initials}".strip()
    return (getattr(user, "full_name", "") or getattr(user, "username", "") or "").strip()


def document_context(name: str, code: str, specialist=None) -> dict:
    today = timezone.localdate()
    specialist_short = short_name(specialist)
    return {
        "organization": getattr(settings, "ORGANIZATION_NAME", ""),
        "city": getattr(settings, "ORGANIZATION_CITY", ""),
        "workshop_name": name,
        "workshop_code": code,
        "date": today.strftime("%d.%m.%Y"),
        "date_long": date_long(today),
        "date_day": f"{today.day:02d}",
        "date_month": MONTHS_RU[today.month - 1],
        "date_year": str(today.year),
        "specialist": specialist_short,
        "specialist_name": specialist_short,
        "specialist_initials": specialist_short,
        "specialist_full_name": getattr(specialist, "full_name", "") if specialist else "",
        "specialist_position": getattr(specialist, "position", "") if specialist else "",
    }


def _render_signature_line(text: str, position: str, specialist: str) -> str:
    """Подпись ``<должность>  _______ <ФИО>`` — должность и ФИО подставляются."""
    match = SIGNATURE_LINE_RE.match(text)
    if not match:
        return text
    before = match.group("before")
    underscores = match.group("us")
    after = match.group("after")
    if position and before.strip():
        trailing = before[len(before.rstrip()):] or "  "
        before = position + trailing
    if specialist and after.strip():
        after = " " + specialist
    return before + underscores + after


def render_sample_text(text: str, context: dict) -> str:
    """Подставляет {{ поля }}, дату, ФИО и подпись специалиста в текст образца."""
    if not text:
        return text
    rendered = render_text(text, context)

    if "«" in rendered:

        def replace(match):
            token = match.group(1)
            if token == "ДАТА":
                return f"«{context.get('date_day', '')}»"
            if token == "МЕСЯЦ":
                return context.get("date_month", "")
            return context.get("date_year", "")

        rendered = DATE_TOKEN_RE.sub(replace, rendered)

    specialist = context.get("specialist")
    position = context.get("specialist_position")
    if specialist:
        rendered = rendered.replace("ФАМИЛИЯ/ИНИЦИАЛЫ", specialist)
        rendered = rendered.replace("ФИО", specialist)
        rendered = FROM_LINE_RE.sub(lambda m: m.group(1) + specialist, rendered)
    if SIGNATURE_LINE_RE.search(rendered):
        rendered = _render_signature_line(rendered, position, specialist)
    return rendered


def _replace_in_paragraph(paragraph, context: dict, state: dict | None = None) -> None:
    if not paragraph.runs:
        return
    original = "".join(run.text for run in paragraph.runs)
    stripped = original.strip()
    rendered = render_sample_text(original, context)

    if state is not None and stripped:
        position = context.get("specialist_position")
        if state.get("after_from"):
            if stripped.startswith("Должность:") and position:
                rendered = POSITION_RE.sub(lambda m: m.group(1) + position, rendered)
            state["after_from"] = False
        if stripped.startswith("От:"):
            state["after_from"] = True

    if rendered == original:
        return
    for run in paragraph.runs:
        run.text = ""
    paragraph.runs[0].text = rendered


def _replace_in_container(container, context: dict) -> None:
    state = {"after_from": False}
    for paragraph in container.paragraphs:
        _replace_in_paragraph(paragraph, context, state)
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                _replace_in_container(cell, context)


def _fill_sample(path: Path, context: dict) -> BytesIO:
    docx = DocxDocument(str(path))
    _replace_in_container(docx, context)
    for section in docx.sections:
        _replace_in_container(section.header, context)
        _replace_in_container(section.footer, context)
    buffer = BytesIO()
    docx.save(buffer)
    buffer.seek(0)
    return buffer


def _build_placeholder(name: str, code: str, kind: str) -> BytesIO:
    docx = DocxDocument()

    style = docx.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    title = docx.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(DOC_KINDS[kind])
    title_run.bold = True
    title_run.font.size = Pt(14)

    docx.add_paragraph()
    docx.add_paragraph(f"Цех: {name} (код {code})")
    docx.add_paragraph(f"Дата: {document_context(name, code)['date']}")
    docx.add_paragraph()

    note = docx.add_paragraph(
        "Образец документа не загружен. Поместите файл "
        f"«{code}_{kind}.docx» в папку documents/samples/, чтобы использовать его как шаблон."
    )
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buffer = BytesIO()
    docx.save(buffer)
    buffer.seek(0)
    return buffer


def build_workshop_document(name: str, code: str, kind: str, specialist=None) -> BytesIO:
    """Возвращает поток DOCX по образцу (или заглушку, если образца нет)."""
    context = document_context(name, code, specialist)
    path = sample_path(code, kind)
    if path.exists():
        return _fill_sample(path, context)
    return _build_placeholder(name, code, kind)

"""Формирование документов: рендер плейсхолдеров и генерация DOCX."""
from __future__ import annotations

from io import BytesIO

from django.conf import settings
from django.core.files.base import ContentFile
from django.utils import timezone
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from documents.models import Document, PLACEHOLDER_RE

MONTHS_RU = [
    "января", "февраля", "марта", "апреля", "мая", "июня",
    "июля", "августа", "сентября", "октября", "ноября", "декабря",
]


def date_long(value) -> str:
    return f"«{value.day:02d}» {MONTHS_RU[value.month - 1]} {value.year} г."


def render_text(text: str, context: dict) -> str:
    def replace(match):
        key = match.group(1)
        return str(context.get(key, match.group(0)))

    return PLACEHOLDER_RE.sub(replace, text or "")


def build_context(document: Document) -> dict:
    workshop = document.workshop
    author = document.created_by
    doc_date = document.doc_date or timezone.localdate()
    context = {
        "organization": getattr(settings, "ORGANIZATION_NAME", ""),
        "city": getattr(settings, "ORGANIZATION_CITY", ""),
        "workshop_name": workshop.name if workshop else "",
        "workshop_code": workshop.code if workshop else "",
        "chief": workshop.chief if workshop else "",
        "chief_position": workshop.chief_position if workshop else "",
        "site": workshop.site.name if workshop and workshop.site else "",
        "date": doc_date.strftime("%d.%m.%Y"),
        "date_long": date_long(doc_date),
        "number": document.number,
        "author": author.full_name if author else "",
        "author_position": getattr(author, "position", "") if author else "",
        "title": document.template.title_template if document.template else "",
    }
    context.update(document.context_data or {})
    return context


def build_docx(document: Document) -> ContentFile:
    docx = DocxDocument()

    section = docx.sections[0]
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    style = docx.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    header = docx.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run(getattr(settings, "ORGANIZATION_NAME", ""))
    run.bold = True

    city_par = docx.add_paragraph()
    city_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    city_par.add_run(getattr(settings, "ORGANIZATION_CITY", ""))

    title_par = docx.add_paragraph()
    title_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_par.add_run(document.title or document.template.name)
    title_run.bold = True
    title_run.font.size = Pt(14)

    meta = docx.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    meta.add_run(f"№ {document.number or '—'} от {document.doc_date:%d.%m.%Y}")

    docx.add_paragraph()

    for block in (document.body or "").split("\n"):
        paragraph = docx.add_paragraph(block)
        paragraph.paragraph_format.first_line_indent = Cm(1.25)
        paragraph.paragraph_format.space_after = Pt(6)

    docx.add_paragraph()
    signature = docx.add_paragraph()
    signature.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    signature.add_run(
        f"{getattr(document.created_by, 'position', '') or 'Сменный специалист'}  "
        f"_______________  {getattr(document.created_by, 'full_name', '') if document.created_by else ''}"
    )

    buffer = BytesIO()
    docx.save(buffer)
    filename = f"document_{document.pk or 'new'}.docx"
    return ContentFile(buffer.getvalue(), name=filename)

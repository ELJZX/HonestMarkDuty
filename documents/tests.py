import tempfile
from io import BytesIO
from pathlib import Path

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT

from accounts.models import User
from core.models import Workshop
from documents.forms import DocumentForm
from documents.models import (
    Document,
    DocumentKind,
    DocumentStatus,
    DocumentTemplate,
    DocumentType,
)
from documents.services import MONTHS_RU, build_context, build_docx, render_text
from documents.workshop_docs import (
    WORKSHOP_DOCUMENTS,
    build_workshop_document,
    document_context,
    render_sample_text,
    short_name,
)
from shifts.models import Shift


class DocumentTemplateModelTests(TestCase):
    def setUp(self):
        self.workshop = Workshop.objects.create(name="Мясной цех", code="МЦ", chief="Кузнецов Д.В.")
        self.template = DocumentTemplate.objects.create(
            name="Заключение",
            doc_type=DocumentType.TECHNICAL_REPORT,
            title_template="Заключение по {{ workshop_name }}",
            body="Начальник {{ chief }}. Причина: {{ reason }}.",
        )

    def test_placeholders_extracted(self):
        self.assertEqual(self.template.placeholders, ["chief", "reason", "workshop_name"])

    def test_available_for_all_when_no_workshops(self):
        self.assertTrue(self.template.available_for(self.workshop))

    def test_available_for_specific_workshop(self):
        other = Workshop.objects.create(name="Другой", code="Д")
        self.template.workshops.set([self.workshop])
        self.assertTrue(self.template.available_for(self.workshop))
        self.assertFalse(self.template.available_for(other))

    def test_str(self):
        self.assertEqual(str(self.template), "Заключение")


class DocumentServiceTests(TestCase):
    def setUp(self):
        self.workshop = Workshop.objects.create(name="Мясной цех", code="МЦ", chief="Кузнецов Д.В.")
        self.template = DocumentTemplate.objects.create(
            name="Заключение",
            title_template="Заключение по {{ workshop_name }}",
            body="Начальник {{ chief }}. Причина: {{ reason }}.",
        )

    def test_render_text(self):
        self.assertEqual(render_text("Привет, {{ name }}!", {"name": "Иван"}), "Привет, Иван!")

    def test_render_keeps_unknown(self):
        self.assertEqual(render_text("{{ unknown }}", {}), "{{ unknown }}")

    def test_document_render_autofills(self):
        document = Document.objects.create(
            template=self.template, workshop=self.workshop, context_data={"reason": "осмотр"}
        )
        document.render()
        self.assertIn("Мясной цех", document.title)
        self.assertIn("Кузнецов Д.В.", document.body)
        self.assertIn("осмотр", document.body)

    def test_build_context(self):
        document = Document.objects.create(template=self.template, workshop=self.workshop)
        context = build_context(document)
        self.assertEqual(context["workshop_code"], "МЦ")
        self.assertIn("organization", context)
        self.assertIn("date_long", context)

    def test_build_docx_non_empty(self):
        document = Document.objects.create(template=self.template, workshop=self.workshop)
        document.render()
        content = build_docx(document)
        self.assertGreater(len(content.read()), 0)


class DocumentFormTests(TestCase):
    def setUp(self):
        self.template = DocumentTemplate.objects.create(
            name="Заключение",
            title_template="Заголовок {{ workshop_name }}",
            body="Причина: {{ reason }}. Дата: {{ date_long }}",
        )

    def test_dynamic_fields_exclude_auto(self):
        form = DocumentForm(initial={"template": self.template.pk})
        self.assertIn("reason", form.dynamic_field_names)
        self.assertNotIn("workshop_name", form.dynamic_field_names)
        self.assertNotIn("date_long", form.dynamic_field_names)

    def test_collect_context(self):
        form = DocumentForm(
            data={"template": self.template.pk, "doc_date": "2026-09-12", "status": "saved", "reason": "тест"}
        )
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.collect_context(), {"reason": "тест"})


class DocumentViewTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            username="admin", password="x", role=User.Role.ADMIN, is_superuser=True, is_staff=True
        )
        self.specialist = User.objects.create_user(
            username="spec", password="x", role=User.Role.SPECIALIST
        )
        self.workshop = Workshop.objects.create(name="Мясной цех", code="МЦ", chief="Кузнецов")
        self.template = DocumentTemplate.objects.create(
            name="Заключение",
            doc_type=DocumentType.TECHNICAL_REPORT,
            title_template="Заключение по {{ workshop_name }}",
            body="Причина: {{ reason }}.",
        )
        self.template.workshops.set([self.workshop])

    def test_template_views(self):
        self.client.force_login(self.specialist)
        self.assertEqual(self.client.get(reverse("documents:template_list")).status_code, 200)
        response = self.client.post(
            reverse("documents:template_create"),
            {
                "name": "Новый шаблон",
                "doc_type": DocumentType.SERVICE_NOTE,
                "title_template": "Записка {{ workshop_name }}",
                "body": "Текст {{ reason }}",
                "is_active": "on",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(DocumentTemplate.objects.filter(name="Новый шаблон").exists())

    def test_document_list_renders(self):
        self.client.force_login(self.specialist)
        self.assertEqual(self.client.get(reverse("documents:document_list")).status_code, 200)

    def test_create_document_generates_docx(self):
        self.client.force_login(self.specialist)
        response = self.client.post(
            reverse("documents:document_create"),
            {
                "template": self.template.pk,
                "workshop": self.workshop.pk,
                "number": "ТЗ-1",
                "doc_date": "2026-09-12",
                "status": DocumentStatus.SAVED,
                "reason": "плановый осмотр",
            },
        )
        self.assertEqual(response.status_code, 302)
        document = Document.objects.get(number="ТЗ-1")
        self.assertIn("Мясной цех", document.title)
        self.assertTrue(document.file)
        self.assertEqual(document.created_by, self.specialist)

    def test_detail_print_download_regenerate(self):
        document = Document.objects.create(
            template=self.template, workshop=self.workshop, number="ТЗ-2", context_data={"reason": "x"}
        )
        document.render()
        document.file.save(f"document_{document.pk}.docx", build_docx(document), save=True)
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse("documents:document_detail", args=[document.pk])).status_code, 200)
        self.assertEqual(self.client.get(reverse("documents:document_print", args=[document.pk])).status_code, 200)
        download = self.client.get(reverse("documents:document_download", args=[document.pk]))
        self.assertEqual(download.status_code, 200)
        self.assertIn("officedocument", download["Content-Type"])
        response = self.client.post(reverse("documents:document_regenerate", args=[document.pk]))
        self.assertEqual(response.status_code, 302)

    def test_update_document(self):
        document = Document.objects.create(
            template=self.template, workshop=self.workshop, number="ТЗ-3", context_data={"reason": "a"}
        )
        document.render()
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse("documents:document_update", args=[document.pk]),
            {
                "template": self.template.pk,
                "workshop": self.workshop.pk,
                "number": "ТЗ-3",
                "doc_date": "2026-09-12",
                "status": DocumentStatus.SAVED,
                "reason": "обновлено",
            },
        )
        self.assertEqual(response.status_code, 302)
        document.refresh_from_db()
        self.assertIn("обновлено", document.body)

    def test_document_list_filters(self):
        Document.objects.create(
            template=self.template,
            workshop=self.workshop,
            number="ТЗ-9",
            title="Заключение",
            status=DocumentStatus.SAVED,
        )
        self.client.force_login(self.specialist)
        response = self.client.get(
            reverse("documents:document_list"),
            {
                "q": "ТЗ-9",
                "type": DocumentType.TECHNICAL_REPORT,
                "workshop": self.workshop.pk,
                "status": DocumentStatus.SAVED,
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "ТЗ-9")

    def test_create_form_preview_with_template(self):
        self.client.force_login(self.specialist)
        response = self.client.get(
            reverse("documents:document_create"),
            {"template": self.template.pk, "workshop": self.workshop.pk},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["preview_title"], "Заключение по Мясной цех")

    def test_download_generates_missing_file(self):
        document = Document.objects.create(
            template=self.template,
            workshop=self.workshop,
            number="ТЗ-4",
            context_data={"reason": "x"},
        )
        document.render()
        self.client.force_login(self.specialist)
        response = self.client.get(reverse("documents:document_download", args=[document.pk]))
        self.assertEqual(response.status_code, 200)
        document.refresh_from_db()
        self.assertTrue(document.file)


class WorkshopDocumentViewTests(TestCase):
    def setUp(self):
        self.specialist = User.objects.create_user(
            username="wspec", password="x", role=User.Role.SPECIALIST
        )

    def test_template_list_lists_workshops_and_buttons(self):
        self.client.force_login(self.specialist)
        response = self.client.get(reverse("documents:template_list"))
        self.assertEqual(response.status_code, 200)
        for item in WORKSHOP_DOCUMENTS:
            self.assertContains(response, item["name"])
        self.assertContains(response, "Скачать техническое заключение")
        self.assertContains(response, "Скачать служебную записку")
        self.assertContains(
            response,
            reverse("documents:workshop_document_download", args=["ceh1", "tz"]),
        )

    def test_download_placeholder_document(self):
        self.client.force_login(self.specialist)
        response = self.client.get(
            reverse("documents:workshop_document_download", args=["mc", "sl"])
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("officedocument", response["Content-Type"])
        expected = f"mc_sl_{timezone.localdate():%d.%m.%Y}.docx"
        self.assertIn(expected, response["Content-Disposition"])
        body = b"".join(response.streaming_content)
        self.assertGreater(len(body), 0)
        text = "\n".join(p.text for p in DocxDocument(BytesIO(body)).paragraphs)
        self.assertIn("Служебная записка", text)
        self.assertIn("Малыш моцарелла", text)

    def test_download_uses_sample_when_present(self):
        with tempfile.TemporaryDirectory() as tmp:
            sample = Path(tmp) / "tv_tz.docx"
            docx = DocxDocument()
            docx.add_paragraph("Цех: {{ workshop_name }} ({{ workshop_code }})")
            docx.add_paragraph("Дата: {{ date }}")
            docx.save(str(sample))

            with override_settings(DOCUMENT_SAMPLES_DIR=tmp):
                self.client.force_login(self.specialist)
                response = self.client.get(
                    reverse("documents:workshop_document_download", args=["tv", "tz"])
                )
                body = b"".join(response.streaming_content)

        text = "\n".join(p.text for p in DocxDocument(BytesIO(body)).paragraphs)
        self.assertIn("Творожный цех", text)
        self.assertIn("(tv)", text)
        self.assertIn(timezone.localdate().strftime("%d.%m.%Y"), text)
        self.assertNotIn("{{", text)

    def test_sample_fills_date_and_specialist(self):
        specialist = User.objects.create_user(
            username="borodin",
            password="x",
            last_name="Бородин",
            first_name="Александр",
            patronymic="Викторович",
            position="Ведущий инженер по цифровой маркировке",
            role=User.Role.SPECIALIST,
        )

        with tempfile.TemporaryDirectory() as tmp:
            sample = Path(tmp) / "mc_tz.docx"
            docx = DocxDocument()
            docx.add_paragraph("«ДАТА» «МЕСЯЦ» «ГОД»")
            docx.add_paragraph("Кому: Начальник")
            docx.add_paragraph("Должность: Начальник цеха №1")
            docx.add_paragraph("От: «ФИО»")
            docx.add_paragraph("Должность: «ДОЛЖНОСТЬ»")
            docx.add_paragraph("От: Пупкин П.П.")
            docx.add_paragraph("«ДОЛЖНОСТЬ»  _______ «ФИО»")
            docx.save(str(sample))

            with override_settings(DOCUMENT_SAMPLES_DIR=tmp):
                self.client.force_login(specialist)
                response = self.client.get(
                    reverse("documents:workshop_document_download", args=["mc", "tz"])
                )
                body = b"".join(response.streaming_content)

        text = "\n".join(p.text for p in DocxDocument(BytesIO(body)).paragraphs)
        today = timezone.localdate()
        self.assertIn(f"«{today.day:02d}» {MONTHS_RU[today.month - 1]} {today.year}", text)
        self.assertIn("От: Бородин А.В.", text)
        self.assertIn("_______ Бородин А.В.", text)
        self.assertIn("Должность: Начальник цеха №1", text)
        self.assertIn("Должность: Ведущий инженер по цифровой маркировке", text)
        self.assertIn("Ведущий инженер по цифровой маркировке\t_______ Бородин А.В.", text)
        self.assertNotIn("Пупкин", text)
        self.assertNotIn("{{", text)
        self.assertNotIn("ДАТА", text)
        self.assertNotIn("«ФИО»", text)
        self.assertNotIn("«ДОЛЖНОСТЬ»", text)

    def test_unknown_code_returns_404(self):
        self.client.force_login(self.specialist)
        response = self.client.get(
            reverse("documents:workshop_document_download", args=["nope", "tz"])
        )
        self.assertEqual(response.status_code, 404)

    def test_requires_login(self):
        response = self.client.get(
            reverse("documents:workshop_document_download", args=["ceh1", "tz"])
        )
        self.assertEqual(response.status_code, 302)


class WorkshopDocumentSampleEdgeTests(TestCase):
    def setUp(self):
        self.specialist = User.objects.create_user(
            username="ws-sample", password="x", role=User.Role.SPECIALIST
        )

    def test_sample_with_plain_paragraph_and_table(self):
        with tempfile.TemporaryDirectory() as tmp:
            sample = Path(tmp) / "tv_tz.docx"
            docx = DocxDocument()
            docx.add_paragraph("Обычный текст без подстановок")
            table = docx.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "Цех: {{ workshop_name }}"
            docx.save(str(sample))

            with override_settings(DOCUMENT_SAMPLES_DIR=tmp):
                self.client.force_login(self.specialist)
                response = self.client.get(
                    reverse("documents:workshop_document_download", args=["tv", "tz"])
                )
                body = b"".join(response.streaming_content)

        parsed = DocxDocument(BytesIO(body))
        paragraphs = "\n".join(p.text for p in parsed.paragraphs)
        self.assertIn("Обычный текст без подстановок", paragraphs)
        table_text = "\n".join(
            cell.text for table in parsed.tables for row in table.rows for cell in row.cells
        )
        self.assertIn("Творожный цех", table_text)


class DocumentFormEdgeTests(TestCase):
    def test_form_without_template_has_no_dynamic_fields(self):
        form = DocumentForm()
        self.assertIsNone(form.template_obj)
        self.assertEqual(form.dynamic_field_names, [])


class DocumentCreatePreviewEdgeTests(TestCase):
    def setUp(self):
        self.specialist = User.objects.create_user(
            username="preview-edge", password="x", role=User.Role.SPECIALIST
        )
        self.workshop = Workshop.objects.create(name="Цех", code="Ц")
        self.template = DocumentTemplate.objects.create(
            name="Шаблон",
            title_template="Т {{ workshop_name }}",
            body="Причина: {{ reason }}",
        )
        self.template.workshops.set([self.workshop])

    def test_invalid_post_builds_preview_with_bad_date(self):
        self.client.force_login(self.specialist)
        response = self.client.post(
            reverse("documents:document_create"),
            {
                "template": self.template.pk,
                "workshop": self.workshop.pk,
                "number": "",
                "doc_date": "не дата",
                "reason": "значение",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("reason", response.context["form"].dynamic_field_names)


class DocumentUploadArchiveTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            username="doc-admin", password="x", role=User.Role.ADMIN, is_superuser=True, is_staff=True
        )
        self.specialist = User.objects.create_user(
            username="doc-spec", password="x", role=User.Role.SPECIALIST
        )
        self.other = User.objects.create_user(
            username="doc-other", password="x", role=User.Role.SPECIALIST
        )
        self.workshop = Workshop.objects.create(name="Цех №1", code="ЦЕХ1")
        self.shift = Shift.objects.create(opened_by=self.specialist, status=Shift.Status.OPEN)

    def _upload(self, name, user=None, content=b"doc-bytes", follow=False):
        self.client.force_login(user or self.specialist)
        return self.client.post(
            reverse("documents:document_upload"),
            {"file": SimpleUploadedFile(name, content)},
            follow=follow,
        )

    @staticmethod
    def _docx_bytes(text):
        buffer = BytesIO()
        docx = DocxDocument()
        docx.add_paragraph(text)
        docx.save(buffer)
        return buffer.getvalue()

    def test_upload_creates_archived_document(self):
        response = self._upload("ceh1_tz_16.09.2026.docx")
        self.assertEqual(response.status_code, 302)
        doc = Document.objects.get()
        self.assertEqual(doc.kind, DocumentKind.TECHNICAL_REPORT)
        self.assertEqual(doc.number, "ТЗ-0001")
        self.assertEqual(doc.workshop, self.workshop)
        self.assertEqual(doc.doc_date.isoformat(), "2026-09-16")
        self.assertEqual(doc.created_by, self.specialist)
        self.assertTrue(doc.file)

    def test_upload_numbers_are_sequential_per_kind(self):
        self._upload("ceh1_tz_16.09.2026.docx")
        self._upload("ceh1_tz_17.09.2026.docx")
        self._upload("csm_sl_17.09.2026.docx")
        numbers = list(Document.objects.order_by("id").values_list("number", flat=True))
        self.assertEqual(numbers, ["ТЗ-0001", "ТЗ-0002", "СЛ-0001"])

    def test_upload_unknown_code_sets_no_workshop(self):
        response = self._upload("zzz_tz_16.09.2026.docx")
        self.assertEqual(response.status_code, 302)
        self.assertIsNone(Document.objects.get().workshop)

    def test_upload_invalid_filename_rejected(self):
        response = self._upload("bad_name.docx")
        self.assertEqual(response.status_code, 302)
        self.assertFalse(Document.objects.exists())

    def test_upload_uses_current_user_without_open_shift(self):
        self.shift.status = Shift.Status.CLOSED
        self.shift.save()
        self._upload("ceh1_tz_16.09.2026.docx", user=self.other)
        self.assertEqual(Document.objects.get().created_by, self.other)

    def test_upload_author_is_logged_in_user(self):
        # Открытая смена принадлежит другому специалисту — автором должен быть загрузивший.
        self._upload(
            "Отчёт 16.09.2026.docx",
            user=self.other,
            content=self._docx_bytes("Техническое заключение"),
        )
        self.assertEqual(Document.objects.get().created_by, self.other)

    def test_upload_detects_kind_from_content(self):
        self._upload(
            "Документ 16.09.2026.docx", content=self._docx_bytes("Техническое заключение")
        )
        doc = Document.objects.get()
        self.assertEqual(doc.kind, DocumentKind.TECHNICAL_REPORT)
        self.assertEqual(doc.number, "ТЗ-0001")

    def test_upload_detects_service_note_from_content(self):
        self._upload(
            "Документ 16.09.2026.docx", content=self._docx_bytes("Служебная записка")
        )
        doc = Document.objects.get()
        self.assertEqual(doc.kind, DocumentKind.SERVICE_NOTE)
        self.assertEqual(doc.number, "СЛ-0001")

    def test_upload_unknown_kind_leaves_blank(self):
        self._upload("Отчёт 16.09.2026.docx", content=self._docx_bytes("Просто текст"))
        doc = Document.objects.get()
        self.assertEqual(doc.kind, "")
        self.assertEqual(doc.number, "")

    def test_upload_parses_flexible_dates(self):
        for name in ("файл_16_09_26.docx", "Отчёт 16.09.26.docx", "файл-16-09-2026.docx"):
            Document.objects.all().delete()
            self._upload(name, content=self._docx_bytes("Техническое заключение"))
            self.assertEqual(Document.objects.get().doc_date.isoformat(), "2026-09-16")

    def test_upload_without_date_shows_error(self):
        response = self._upload(
            "Отчёт без даты.docx",
            content=self._docx_bytes("Текст"),
            follow=True,
        )
        self.assertFalse(Document.objects.exists())
        self.assertContains(response, "В названии файла должна быть указана дата в формате")

    def test_viewer_cannot_upload(self):
        viewer = User.objects.create_user(
            username="doc-viewer", password="x", role=User.Role.VIEWER
        )
        self.client.force_login(viewer)
        response = self.client.post(
            reverse("documents:document_upload"),
            {"file": SimpleUploadedFile("ceh1_tz_16.09.2026.docx", b"x")},
        )
        self.assertEqual(response.status_code, 403)


class DocumentUploadHelpersTests(TestCase):
    def test_parse_date_from_name_formats(self):
        from documents.views import parse_date_from_name

        for stem in ("ceh1_tz_16.09.2026", "файл_16_09_26", "Отчёт 16.09.26", "файл-16-09-2026"):
            self.assertEqual(parse_date_from_name(stem).isoformat(), "2026-09-16", stem)

    def test_parse_date_from_name_rejects_invalid(self):
        from documents.views import parse_date_from_name

        self.assertIsNone(parse_date_from_name("Отчёт без даты"))
        self.assertIsNone(parse_date_from_name("99.99.2026"))

    def test_kind_from_name(self):
        from documents.views import kind_from_name

        self.assertEqual(kind_from_name("ceh1_tz_16.09.2026"), DocumentKind.TECHNICAL_REPORT)
        self.assertEqual(kind_from_name("csm_sl_16.09.2026"), DocumentKind.SERVICE_NOTE)
        self.assertEqual(kind_from_name("просто файл 16.09.2026"), "")

    def test_workshop_from_name(self):
        from documents.views import workshop_from_name

        workshop = Workshop.objects.create(name="Цех №1", code="ЦЕХ1")
        self.assertEqual(workshop_from_name("ceh1_tz_16.09.2026"), workshop)
        self.assertIsNone(workshop_from_name("zzz_tz_16.09.2026"))


class DocumentArchiveSearchTests(TestCase):
    def setUp(self):
        self.specialist = User.objects.create_user(
            username="arch-spec",
            password="x",
            role=User.Role.SPECIALIST,
            last_name="Иванов",
            first_name="Иван",
        )
        self.workshop = Workshop.objects.create(name="Цех №1", code="ЦЕХ1")
        self.document = Document.objects.create(
            kind=DocumentKind.TECHNICAL_REPORT,
            workshop=self.workshop,
            number="ТЗ-0001",
            doc_date=timezone.localdate().replace(year=2026),
            created_by=self.specialist,
        )

    def test_page_shows_archive_and_upload_button(self):
        self.client.force_login(self.specialist)
        response = self.client.get(reverse("documents:document_list"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Архив технических заключений и служебных записок")
        self.assertContains(response, "Добавить документ в архив")
        self.assertContains(response, 'id="doc-search"')
        self.assertNotContains(response, 'placeholder="заголовок')

    def test_partial_returns_only_rows(self):
        self.client.force_login(self.specialist)
        response = self.client.get(reverse("documents:document_list"), {"partial": "1"})
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "<!DOCTYPE")
        self.assertContains(response, "ТЗ-0001")

    def test_search_by_number_author_workshop_date(self):
        self.client.force_login(self.specialist)
        for query in ("ТЗ-0001", "Иванов", "Цех №1", "2026"):
            response = self.client.get(
                reverse("documents:document_list"), {"q": query, "partial": "1"}
            )
            self.assertContains(response, "ТЗ-0001", msg_prefix=query)

    def test_filter_by_kind_and_workshop(self):
        Document.objects.create(
            kind=DocumentKind.SERVICE_NOTE, number="СЛ-0001", workshop=self.workshop
        )
        self.client.force_login(self.specialist)
        response = self.client.get(
            reverse("documents:document_list"),
            {"kind": DocumentKind.SERVICE_NOTE, "partial": "1"},
        )
        self.assertContains(response, "СЛ-0001")
        self.assertNotContains(response, "ТЗ-0001")

        response = self.client.get(
            reverse("documents:document_list"), {"workshop": "Цех №1", "partial": "1"}
        )
        self.assertContains(response, "ТЗ-0001")

    def test_search_no_match(self):
        self.client.force_login(self.specialist)
        response = self.client.get(
            reverse("documents:document_list"), {"q": "нет-такого", "partial": "1"}
        )
        self.assertContains(response, "Документов нет")


class DocumentArchiveActionsTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            username="arch-admin", password="x", role=User.Role.ADMIN, is_superuser=True, is_staff=True
        )
        self.specialist = User.objects.create_user(
            username="arch-user", password="x", role=User.Role.SPECIALIST
        )
        self.document = Document.objects.create(
            kind=DocumentKind.TECHNICAL_REPORT,
            number="ТЗ-0001",
            doc_date=timezone.localdate(),
            created_by=self.specialist,
        )

    def test_no_open_button_in_archive(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse("documents:document_list"))
        self.assertNotContains(response, "Открыть")
        self.assertNotContains(
            response,
            f'href="{reverse("documents:document_detail", args=[self.document.pk])}"',
        )

    def test_delete_button_visible_only_for_admin(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse("documents:document_list"))
        self.assertContains(
            response, reverse("documents:document_delete", args=[self.document.pk])
        )

        self.client.force_login(self.specialist)
        response = self.client.get(reverse("documents:document_list"))
        self.assertNotContains(
            response, reverse("documents:document_delete", args=[self.document.pk])
        )

    def test_admin_can_delete_document(self):
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse("documents:document_delete", args=[self.document.pk])
        )
        self.assertEqual(response.status_code, 302)
        self.assertFalse(Document.objects.filter(pk=self.document.pk).exists())

    def test_specialist_cannot_delete_document(self):
        self.client.force_login(self.specialist)
        response = self.client.post(
            reverse("documents:document_delete", args=[self.document.pk])
        )
        self.assertEqual(response.status_code, 403)
        self.assertTrue(Document.objects.filter(pk=self.document.pk).exists())


class DocumentTemplateListCleanupTests(TestCase):
    def setUp(self):
        self.specialist = User.objects.create_user(
            username="tmpl-spec", password="x", role=User.Role.SPECIALIST
        )

    def test_templates_section_removed(self):
        self.client.force_login(self.specialist)
        response = self.client.get(reverse("documents:template_list"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response, "Готовые шаблоны технических заключений и служебных записок по цехам"
        )
        self.assertNotContains(response, "Шаблоны документов")
        self.assertNotContains(response, "Добавить шаблон")
        self.assertNotContains(response, "+ Шаблон")


class WorkshopDocumentDateTests(TestCase):
    def setUp(self):
        self.specialist = User.objects.create_user(
            username="date-spec", password="x", role=User.Role.SPECIALIST
        )

    def test_date_tokens_replaced(self):
        from documents.workshop_docs import render_sample_text

        today = timezone.localdate()
        rendered = render_sample_text(
            "Дата: «ДАТА» «МЕСЯЦ» «ГОД»",
            {
                "date_day": f"{today.day:02d}",
                "date_month": "сентября",
                "date_year": str(today.year),
            },
        )
        self.assertEqual(rendered, f"Дата: «{today.day:02d}» сентября {today.year}")

    def test_sample_with_date_tokens_downloaded_filled(self):
        with tempfile.TemporaryDirectory() as tmp:
            sample = Path(tmp) / "ceh1_tz.docx"
            docx = DocxDocument()
            docx.add_paragraph("ПАО Комбинат   «ДАТА» «МЕСЯЦ» «ГОД»")
            docx.add_paragraph("Цех: {{ workshop_name }}")
            docx.save(str(sample))

            with override_settings(DOCUMENT_SAMPLES_DIR=tmp):
                self.client.force_login(self.specialist)
                response = self.client.get(
                    reverse("documents:workshop_document_download", args=["ceh1", "tz"])
                )
                body = b"".join(response.streaming_content)

        text = "\n".join(p.text for p in DocxDocument(BytesIO(body)).paragraphs)
        today = timezone.localdate()
        self.assertIn(f"«{today.day:02d}»", text)
        self.assertIn(str(today.year), text)
        self.assertIn("Цех №1", text)
        self.assertNotIn("«ДАТА»", text)


class WorkshopDocsFunctionTests(TestCase):
    """Юнит-тесты функций формирования документов по цехам."""

    def _user(self, **overrides):
        data = {
            "username": "doc-user",
            "password": "x",
            "last_name": "Бобров",
            "first_name": "Максим",
            "patronymic": "Андреевич",
            "position": "Ведущий инженер по цифровой маркировке",
            "role": User.Role.SPECIALIST,
        }
        data.update(overrides)
        return User.objects.create_user(**data)

    def test_short_name_full(self):
        self.assertEqual(short_name(self._user()), "Бобров М.А.")

    def test_short_name_without_last_name_falls_back(self):
        user = self._user(username="fallback", last_name="", first_name="", patronymic="")
        self.assertEqual(short_name(user), "fallback")

    def test_short_name_none(self):
        self.assertEqual(short_name(None), "")

    def test_document_context_fields(self):
        context = document_context("Цех №1", "ceh1", self._user())
        self.assertEqual(context["workshop_name"], "Цех №1")
        self.assertEqual(context["workshop_code"], "ceh1")
        self.assertEqual(context["specialist"], "Бобров М.А.")
        self.assertEqual(
            context["specialist_position"], "Ведущий инженер по цифровой маркировке"
        )
        self.assertEqual(context["date_day"], f"{timezone.localdate().day:02d}")

    def test_render_sample_text_fills_from_line(self):
        context = document_context("Цех №1", "ceh1", self._user())
        self.assertEqual(render_sample_text("От: «ФИО»", context), "От: Бобров М.А.")
        self.assertEqual(render_sample_text("От: Иванов И.И.", context), "От: Бобров М.А.")

    def test_render_sample_text_fills_signature_position_and_name(self):
        context = document_context("Цех №1", "ceh1", self._user())
        rendered = render_sample_text(
            "Специалист по маркировке  _______ «ФАМИЛИЯ/ИНИЦИАЛЫ»", context
        )
        self.assertEqual(
            rendered, "Ведущий инженер по цифровой маркировке\t_______ Бобров М.А."
        )

    def test_render_sample_text_fills_quoted_position(self):
        context = document_context("Цех №1", "ceh1", self._user())
        self.assertEqual(
            render_sample_text("Должность: «ДОЛЖНОСТЬ»", context),
            "Должность: Ведущий инженер по цифровой маркировке",
        )

    def test_render_sample_text_signature_with_two_underscore_groups(self):
        context = document_context("Цех №1", "ceh1", self._user())
        self.assertEqual(
            render_sample_text("«ДОЛЖНОСТЬ» ___     _____ «ФИО»", context),
            "Ведущий инженер по цифровой маркировке\t___     _____ Бобров М.А.",
        )

    def test_render_sample_text_keeps_chief_position(self):
        context = document_context("Цех №1", "ceh1", self._user())
        self.assertEqual(
            render_sample_text("Должность: Начальник цеха №1", context),
            "Должность: Начальник цеха №1",
        )

    def test_build_workshop_document_fills_specialist_position_after_from(self):
        with tempfile.TemporaryDirectory() as tmp:
            sample = Path(tmp) / "ceh1_tz.docx"
            docx = DocxDocument()
            docx.add_paragraph("Кому: Начальник")
            docx.add_paragraph("Должность: Начальник цеха №1")
            docx.add_paragraph("От: Иванов И.И.")
            docx.add_paragraph("Должность: Специалист автоматизированных систем маркировки")
            docx.save(str(sample))

            with override_settings(DOCUMENT_SAMPLES_DIR=tmp):
                stream = build_workshop_document("Цех №1", "ceh1", "tz", self._user())

        texts = [p.text for p in DocxDocument(stream).paragraphs]
        self.assertIn("Должность: Начальник цеха №1", texts)
        self.assertIn("Должность: Ведущий инженер по цифровой маркировке", texts)
        self.assertIn("От: Бобров М.А.", texts)

    def test_signature_line_uses_right_tab_stop(self):
        with tempfile.TemporaryDirectory() as tmp:
            sample = Path(tmp) / "ceh1_tz.docx"
            docx = DocxDocument()
            docx.add_paragraph("Обычный абзац без подписи")
            docx.add_paragraph("Ведущий инженер  _______ «ФИО»")
            docx.save(str(sample))

            with override_settings(DOCUMENT_SAMPLES_DIR=tmp):
                stream = build_workshop_document("Цех №1", "ceh1", "tz", self._user())

        paragraphs = list(DocxDocument(stream).paragraphs)
        self.assertIsNone(paragraphs[0].alignment)

        signature = paragraphs[1]
        self.assertEqual(signature.alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertIn("\t", signature.text)
        self.assertTrue(
            any(
                stop.alignment == WD_TAB_ALIGNMENT.RIGHT
                for stop in signature.paragraph_format.tab_stops
            )
        )

    def test_build_workshop_document_placeholder_without_sample(self):
        with tempfile.TemporaryDirectory() as tmp:
            with override_settings(DOCUMENT_SAMPLES_DIR=tmp):
                stream = build_workshop_document("Малыш сырки", "ms", "sl", self._user())
        text = "\n".join(p.text for p in DocxDocument(stream).paragraphs)
        self.assertIn("Служебная записка", text)
        self.assertIn("Малыш сырки", text)

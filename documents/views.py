from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin
from django.db.models import CharField, Q
from django.db.models.functions import Cast
from django.http import FileResponse, Http404
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse_lazy
from django.utils import timezone
from django.views.generic import (
    CreateView,
    DeleteView,
    DetailView,
    ListView,
    UpdateView,
    View,
)

from docx import Document as DocxDocument

from core.mixins import AdminRequiredMixin, EditorRequiredMixin
from core.models import Workshop
from documents.forms import DocumentForm, DocumentTemplateForm
from documents.models import (
    Document,
    DocumentKind,
    DocumentTemplate,
    DocumentType,
)
from documents.services import build_context, build_docx, render_text
from documents.workshop_docs import (
    DOC_KINDS,
    WORKSHOP_BY_CODE,
    WORKSHOP_DOCUMENTS,
    build_workshop_document,
)


class DocumentTemplateListView(LoginRequiredMixin, ListView):
    model = DocumentTemplate
    template_name = "documents/template_list.html"
    context_object_name = "templates"

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["workshop_rows"] = WORKSHOP_DOCUMENTS
        return ctx


class WorkshopDocumentDownloadView(LoginRequiredMixin, View):
    """Скачивание документа цеха (техническое заключение / служебная записка)."""

    DOCX_CONTENT_TYPE = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    def get(self, request, code, kind):
        name = WORKSHOP_BY_CODE.get(code)
        if name is None or kind not in DOC_KINDS:
            raise Http404("Документ не найден")

        stream = build_workshop_document(name, code, kind, specialist=request.user)
        filename = f"{code}_{kind}_{timezone.localdate():%d.%m.%Y}.docx"
        return FileResponse(
            stream,
            as_attachment=True,
            filename=filename,
            content_type=self.DOCX_CONTENT_TYPE,
        )


class DocumentTemplateCreateView(EditorRequiredMixin, CreateView):
    model = DocumentTemplate
    form_class = DocumentTemplateForm
    template_name = "documents/template_form.html"
    success_url = reverse_lazy("documents:template_list")
    extra_context = {"title": "Новый шаблон документа", "back_url": "documents:template_list"}

    def form_valid(self, form):
        messages.success(self.request, "Шаблон создан.")
        return super().form_valid(form)


class DocumentTemplateUpdateView(AdminRequiredMixin, UpdateView):
    model = DocumentTemplate
    form_class = DocumentTemplateForm
    template_name = "documents/template_form.html"
    success_url = reverse_lazy("documents:template_list")
    extra_context = {"title": "Редактирование шаблона", "back_url": "documents:template_list"}


class DocumentTemplateDeleteView(AdminRequiredMixin, DeleteView):
    model = DocumentTemplate
    template_name = "core/confirm_delete.html"
    success_url = reverse_lazy("documents:template_list")


DATE_IN_NAME_RE = re.compile(
    r"(?<!\d)(?P<day>\d{1,2})[._/\-](?P<month>\d{1,2})[._/\-](?P<year>\d{2,4})(?!\d)"
)
KIND_TOKEN_RE = re.compile(r"(?<![a-z0-9])(?P<kind>tz|sl)(?![a-z0-9])", re.IGNORECASE)
DOCUMENT_NUMBER_RE = re.compile(r"^[А-Яа-яA-Za-z]{2}-(\d+)$")

DATE_ERROR_MESSAGE = (
    "В названии файла должна быть указана дата в формате — ДД.ММ.ГГГГ, ДД.ММ.ГГ, "
    "ДД_ММ_ГГГГ, ДД/ММ/ГГГГ (например: 16.09.2026, 16_09_26, 16/09/2026)."
)


def parse_date_from_name(stem: str):
    """Ищет дату в имени файла: 16.09.26, 16_09_26, 16/09/2026 и т.п."""
    match = DATE_IN_NAME_RE.search(stem or "")
    if not match:
        return None
    day = int(match.group("day"))
    month = int(match.group("month"))
    year = int(match.group("year"))
    if year < 100:
        year += 2000
    try:
        return date(year, month, day)
    except ValueError:
        return None


def kind_from_name(stem: str) -> str:
    """Вид документа по токену tz/sl в имени файла (если есть)."""
    match = KIND_TOKEN_RE.search(stem or "")
    return match.group("kind").lower() if match else ""


def kind_from_content(upload) -> str:
    """Вид документа по тексту внутри файла."""
    name = (upload.name or "").lower()
    text = ""
    if name.endswith(".docx"):
        try:
            upload.seek(0)
            docx = DocxDocument(upload)
            chunks = [paragraph.text for paragraph in docx.paragraphs]
            for table in docx.tables:
                for row in table.rows:
                    for cell in row.cells:
                        chunks.append(cell.text)
            text = "\n".join(chunks)
        except Exception:  # noqa: BLE001
            text = ""
    if not text:
        try:
            upload.seek(0)
            text = upload.read().decode("utf-8", errors="ignore")
        except Exception:  # noqa: BLE001
            text = ""
    lowered = text.lower()
    if "служебная записка" in lowered:
        return DocumentKind.SERVICE_NOTE
    if "техническое заключение" in lowered:
        return DocumentKind.TECHNICAL_REPORT
    return ""


def workshop_from_name(stem: str):
    """Цех по коду (ceh1, kmc, …) в имени файла, если он там есть."""
    lowered = (stem or "").lower()
    for code, name in WORKSHOP_BY_CODE.items():
        if re.search(rf"(?<![a-z0-9]){re.escape(code)}(?![a-z0-9])", lowered):
            return Workshop.objects.filter(name=name).first()
    return None


def next_document_number(kind: str) -> str:
    """Следующий номер документа: ТЗ-0001 / СЛ-0001."""
    prefix = "ТЗ" if kind == DocumentKind.TECHNICAL_REPORT else "СЛ"
    max_number = 0
    for number in Document.objects.filter(kind=kind).values_list("number", flat=True):
        match = DOCUMENT_NUMBER_RE.match(number or "")
        if match:
            max_number = max(max_number, int(match.group(1)))
    return f"{prefix}-{max_number + 1:04d}"


def kind_from_template(template) -> str:
    if not template:
        return ""
    mapping = {
        DocumentType.TECHNICAL_REPORT: DocumentKind.TECHNICAL_REPORT,
        DocumentType.SERVICE_NOTE: DocumentKind.SERVICE_NOTE,
    }
    return mapping.get(template.doc_type, "")


class DocumentListView(LoginRequiredMixin, ListView):
    model = Document
    template_name = "documents/document_list.html"
    context_object_name = "documents"
    paginate_by = 25

    def get_queryset(self):
        qs = Document.objects.select_related("template", "workshop", "created_by")
        params = self.request.GET
        query = (params.get("q") or "").strip()
        if query:
            qs = qs.annotate(
                doc_date_text=Cast("doc_date", output_field=CharField())
            ).filter(
                Q(number__icontains=query)
                | Q(title__icontains=query)
                | Q(workshop__name__icontains=query)
                | Q(created_by__last_name__icontains=query)
                | Q(created_by__first_name__icontains=query)
                | Q(created_by__patronymic__icontains=query)
                | Q(created_by__username__icontains=query)
                | Q(doc_date_text__icontains=query)
            )
        if params.get("kind"):
            qs = qs.filter(kind=params["kind"])
        if params.get("workshop"):
            qs = qs.filter(workshop__name=params["workshop"])
        return qs

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["kinds"] = DocumentKind.choices
        ctx["workshop_rows"] = WORKSHOP_DOCUMENTS
        ctx["current"] = self.request.GET
        return ctx

    def render_to_response(self, context, **response_kwargs):
        if self.request.GET.get("partial"):
            return render(self.request, "documents/partials/document_rows.html", context)
        return super().render_to_response(context, **response_kwargs)


class DocumentUploadView(EditorRequiredMixin, View):
    """Добавление готового документа в архив. Дата берётся из имени файла."""

    def post(self, request):
        upload = request.FILES.get("file")
        if upload is None:
            messages.error(request, "Файл не выбран.")
            return redirect("documents:document_list")

        stem = Path(upload.name).stem
        doc_date = parse_date_from_name(stem)
        if doc_date is None:
            messages.error(request, DATE_ERROR_MESSAGE)
            return redirect("documents:document_list")

        kind = kind_from_content(upload) or kind_from_name(stem)

        document = Document(
            kind=kind,
            workshop=workshop_from_name(stem),
            doc_date=doc_date,
            number=next_document_number(kind) if kind else "",
            created_by=request.user,
        )
        upload.seek(0)
        document.file.save(Path(upload.name).name, upload, save=False)
        document.save()
        if document.number:
            messages.success(request, f"Документ {document.number} добавлен в архив.")
        else:
            messages.success(request, "Документ добавлен в архив.")
        return redirect("documents:document_list")


class DocumentCreateView(EditorRequiredMixin, CreateView):
    model = Document
    form_class = DocumentForm
    template_name = "documents/document_form.html"

    def get_initial(self):
        initial = super().get_initial()
        if self.request.GET.get("template"):
            initial["template"] = self.request.GET["template"]
        if self.request.GET.get("workshop"):
            initial["workshop"] = self.request.GET["workshop"]
        return initial

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        form = ctx["form"]
        ctx["title"] = "Новый документ"
        ctx["back_url"] = "documents:document_list"
        ctx["preview_title"] = ""
        ctx["preview_body"] = ""
        template = getattr(form, "template_obj", None)
        if template:
            from datetime import date as date_cls

            from django.utils import timezone

            raw_date = form["doc_date"].value()
            doc_date = None
            if raw_date:
                try:
                    doc_date = date_cls.fromisoformat(str(raw_date))
                except ValueError:
                    doc_date = None
            draft = Document(
                template=template,
                workshop=self._workshop(form),
                number=form["number"].value() or "",
                doc_date=doc_date or timezone.localdate(),
                created_by=self.request.user,
                context_data=self._posted_context(form),
            )
            context = build_context(draft)
            ctx["preview_title"] = render_text(template.title_template, context)
            ctx["preview_body"] = render_text(template.body, context)
        return ctx

    @staticmethod
    def _workshop(form):
        from core.models import Workshop

        value = form["workshop"].value()
        return Workshop.objects.filter(pk=value).first() if value else None

    @staticmethod
    def _posted_context(form):
        data = {}
        for name in getattr(form, "dynamic_field_names", []):
            value = form[name].value()
            if value:
                data[name] = value
        return data

    def form_valid(self, form):
        document = form.save(commit=False)
        document.created_by = self.request.user
        document.kind = kind_from_template(document.template)
        document.context_data = form.collect_context()
        document.save()
        document.render(save=False)
        document.file.save(f"document_{document.pk}.docx", build_docx(document), save=False)
        document.save()
        messages.success(self.request, "Документ сформирован.")
        return redirect("documents:document_detail", pk=document.pk)


class DocumentUpdateView(AdminRequiredMixin, UpdateView):
    model = Document
    form_class = DocumentForm
    template_name = "documents/document_form.html"
    extra_context = {"title": "Редактирование документа", "back_url": "documents:document_list"}

    def form_valid(self, form):
        document = form.save(commit=False)
        document.kind = kind_from_template(document.template)
        document.context_data = form.collect_context()
        document.save()
        document.render(save=False)
        document.file.save(f"document_{document.pk}.docx", build_docx(document), save=False)
        document.save()
        messages.success(self.request, "Документ обновлён и пересформирован.")
        return redirect("documents:document_detail", pk=document.pk)


class DocumentDetailView(LoginRequiredMixin, DetailView):
    model = Document
    template_name = "documents/document_detail.html"
    context_object_name = "document"

    def get_queryset(self):
        return Document.objects.select_related("template", "workshop", "created_by")


class DocumentPrintView(LoginRequiredMixin, DetailView):
    model = Document
    template_name = "documents/document_print.html"
    context_object_name = "document"


class DocumentDownloadView(LoginRequiredMixin, View):
    DOCX_CONTENT_TYPE = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    def get(self, request, pk):
        document = get_object_or_404(Document, pk=pk)
        if not document.file and document.template:
            document.file.save(f"document_{document.pk}.docx", build_docx(document), save=True)
        if not document.file:
            raise Http404("Файл документа не найден.")

        extension = Path(document.file.name).suffix.lower() or ".docx"
        content_types = {
            ".docx": self.DOCX_CONTENT_TYPE,
            ".doc": "application/msword",
        }
        filename = f"{document.kind or 'document'}_{document.number or document.pk}{extension}".replace(" ", "_")
        return FileResponse(
            document.file.open("rb"),
            as_attachment=True,
            filename=filename,
            content_type=content_types.get(extension, "application/octet-stream"),
        )


class DocumentRegenerateView(AdminRequiredMixin, View):
    def post(self, request, pk):
        document = get_object_or_404(Document, pk=pk)
        document.render(save=False)
        document.file.save(f"document_{document.pk}.docx", build_docx(document), save=False)
        document.save()
        messages.success(request, "Документ пересформирован.")
        return redirect("documents:document_detail", pk=pk)


class DocumentDeleteView(AdminRequiredMixin, DeleteView):
    model = Document
    template_name = "core/confirm_delete.html"
    success_url = reverse_lazy("documents:document_list")

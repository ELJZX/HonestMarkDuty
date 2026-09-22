"""Синхронизация смен и журнала с внешней системой производства.

Источник данных — существующий сайт производства (Django). Мы только читаем:
- выгрузка событий смены ``/main/shift/export_events/`` (XLSX: дата/время, ФИО,
  тип события, комментарий);
- таблица графика ``/main/shift/`` (День, ФИО, начало, конец).

Специалист принимает/сдаёт смену там, мы лишь отражаем факт у себя.
"""
from __future__ import annotations

import hashlib
import http.cookiejar
import io
import re
import urllib.parse
import urllib.request
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta

from django.conf import settings
from django.core.cache import cache
from django.utils import timezone

MONTHS = {
    "января": 1, "февраля": 2, "марта": 3, "апреля": 4, "мая": 5, "июня": 6,
    "июля": 7, "августа": 8, "сентября": 9, "октября": 10, "ноября": 11, "декабря": 12,
}

_TRANSLIT = {
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e", "ж": "zh",
    "з": "z", "и": "i", "й": "y", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o",
    "п": "p", "р": "r", "с": "s", "т": "t", "у": "u", "ф": "f", "х": "h", "ц": "ts",
    "ч": "ch", "ш": "sh", "щ": "sch", "ъ": "", "ы": "y", "ь": "", "э": "e", "ю": "yu",
    "я": "ya",
}


@dataclass
class ExternalEvent:
    occurred_at: datetime
    fio: str
    kind: str  # start | end | work
    comment: str


@dataclass
class ExternalWorkday:
    date: date
    fio: str
    start: time
    end: time


def _translit(value: str) -> str:
    return "".join(_TRANSLIT.get(ch.lower(), ch.lower()) for ch in value if ch.isalnum())


def _strip_tags(html: str) -> str:
    return re.sub(r"<[^>]+>", "", html).strip()


def _parse_date_ru(value: str) -> date | None:
    match = re.search(r"(\d{1,2})\s+([А-Яа-яЁё]+)\s+(\d{4})", value)
    if not match:
        return None
    day, month_name, year = match.groups()
    month = MONTHS.get(month_name.lower())
    if not month:
        return None
    return date(int(year), month, int(day))


def _parse_time(value: str) -> time | None:
    match = re.match(r"\s*(\d{1,2}):(\d{2})", value or "")
    if not match:
        return None
    hour, minute = int(match.group(1)), int(match.group(2))
    if hour > 23 or minute > 59:
        return None
    return time(hour, minute)


def _parse_dt(value: str) -> datetime | None:
    for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M"):
        try:
            naive = datetime.strptime(value, fmt)
        except ValueError:
            continue
        return timezone.make_aware(naive, timezone.get_current_timezone())
    return None


# --------------------------------------------------------------------------- HTTP
def login(base_url: str | None = None) -> urllib.request.OpenerDirector:
    base = (base_url or settings.SHIFT_SYNC_URL).rstrip("/")
    jar = http.cookiejar.CookieJar()
    opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))
    login_page = opener.open(f"{base}/main/login/", timeout=20).read().decode("utf-8", "replace")
    token = re.search(r'name="csrfmiddlewaretoken" value="([^"]+)"', login_page)
    payload = urllib.parse.urlencode(
        {
            "username": settings.SHIFT_SYNC_USER,
            "password": settings.SHIFT_SYNC_PASSWORD,
            "csrfmiddlewaretoken": token.group(1) if token else "",
        }
    ).encode()
    request = urllib.request.Request(
        f"{base}/main/login/",
        data=payload,
        method="POST",
        headers={"Referer": f"{base}/main/login/"},
    )
    opener.open(request, timeout=20).read()
    return opener


def fetch_events_xlsx(opener, date_from: date, date_to: date, base_url: str | None = None) -> bytes:
    base = (base_url or settings.SHIFT_SYNC_URL).rstrip("/")
    query = urllib.parse.urlencode(
        {"date_start": date_from.isoformat(), "date_end": date_to.isoformat()}
    )
    with opener.open(f"{base}/main/shift/export_events/?{query}", timeout=60) as response:
        return response.read()


def fetch_schedule_html(opener, base_url: str | None = None) -> str:
    base = (base_url or settings.SHIFT_SYNC_URL).rstrip("/")
    with opener.open(f"{base}/main/shift/", timeout=30) as response:
        return response.read().decode("utf-8", "replace")


def fetch_current_duty(opener, base_url: str | None = None) -> tuple[bool, str]:
    """Один запрос к главной: кто сейчас на смене и работает ли он."""
    base = (base_url or settings.SHIFT_SYNC_URL).rstrip("/")
    with opener.open(f"{base}/main/", timeout=15) as response:
        html = response.read().decode("utf-8", "replace")
    on_duty = "duty-on-work" in html
    match = re.search(r'current-duty-chz-fio[^>]*>([^<]*)<', html)
    fio = (match.group(1) if match else "").strip().rstrip(":").strip()
    return on_duty, fio


# ------------------------------------------------------------------------ parse
def parse_events_xlsx(content: bytes) -> list[ExternalEvent]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    known = {"Начало смены": "start", "Конец смены": "end", "Запись": "work"}
    events: list[ExternalEvent] = []
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        dt = _parse_dt(str(row[0]).strip())
        if dt is None:
            continue
        fio = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        kind = known.get(str(row[2]).strip() if len(row) > 2 and row[2] else "")
        if not kind:
            continue
        comment = str(row[3]) if len(row) > 3 and row[3] is not None else ""
        events.append(ExternalEvent(occurred_at=dt, fio=fio, kind=kind, comment=comment))
    return events


def parse_schedule_html(html: str) -> list[ExternalWorkday]:
    result: list[ExternalWorkday] = []
    for row_html in re.findall(r"<tr[^>]*>(.*?)</tr>", html, re.S | re.I):
        cells = [
            _strip_tags(cell)
            for cell in re.findall(r"<td[^>]*>(.*?)</td>", row_html, re.S | re.I)
        ]
        if len(cells) < 4:
            continue
        day = _parse_date_ru(cells[0])
        start = _parse_time(cells[2])
        end = _parse_time(cells[3])
        if not day or not start:
            continue
        result.append(ExternalWorkday(date=day, fio=cells[1], start=start, end=end or time(20, 0)))
    return result


def parse_work_comment(comment: str) -> dict:
    """Разбирает комментарий события «Запись» (табы: время, линия, ..., действие, решение)."""
    parts = [part.strip() for part in comment.split("\t")]
    while parts and parts[-1] == "":
        parts.pop()
    data = {"equipment_line": "", "downtime": "", "action_task": "", "solution": ""}
    if len(parts) >= 2:
        data["equipment_line"] = parts[1]
    if len(parts) >= 4:
        data["action_task"] = parts[3]
    elif len(parts) == 3:
        data["action_task"] = parts[2]
    if len(parts) >= 5:
        data["solution"] = parts[4]
    if parts:
        data["downtime"] = parts[0]
    return data


# ------------------------------------------------------------------------- users
def user_for_fio(fio: str):
    from accounts.models import User

    fio = (fio or "").strip()
    if not fio:
        return None
    parts = fio.split()
    last = parts[0]
    first = parts[1] if len(parts) > 1 else ""

    user = User.objects.filter(last_name=last, first_name=first).first()
    if user:
        return user
    for candidate in User.objects.exclude(last_name="").exclude(first_name=""):
        if candidate.full_name == fio:
            return candidate
    if not settings.SHIFT_SYNC_CREATE_USERS:
        return None
    username = (_translit(last) + "_" + _translit(first))[:20] or "employee"
    base = username
    index = 1
    while User.objects.filter(username=username).exists():
        index += 1
        username = f"{base}{index}"
    return User.objects.create(
        username=username,
        last_name=last,
        first_name=first,
        patronymic=parts[2] if len(parts) > 2 else "",
        role=User.Role.SPECIALIST,
    )


# -------------------------------------------------------------------------- sync
def _shift_for_date(day: date, fio: str = ""):
    from shifts.models import Shift

    external_id = f"workday:{day.isoformat()}"
    shift = Shift.objects.filter(external_id=external_id).first()
    if shift:
        return shift
    tz = timezone.get_current_timezone()
    today = timezone.localdate()
    opened_at = timezone.make_aware(datetime.combine(day, time(8, 0)), tz)
    if day > today:
        status, closed_at, accepted = Shift.Status.PLANNED, None, False
    elif day < today:
        status = Shift.Status.CLOSED
        closed_at = timezone.make_aware(datetime.combine(day, time(20, 0)), tz)
        accepted = True
    else:
        status, closed_at, accepted = Shift.Status.PLANNED, None, False
    return Shift.objects.create(
        external_id=external_id,
        date=day,
        kind=Shift.Kind.DAY,
        opened_by=user_for_fio(fio),
        opened_at=opened_at,
        closed_at=closed_at,
        status=status,
        accepted=accepted,
    )


def sync_window(days_back: int, days_ahead: int, opener=None, force: bool = False) -> dict:
    from journal.models import JournalEntry
    from shifts.models import Shift

    if not settings.SHIFT_SYNC_ENABLED and not force:
        return {"skipped": "SHIFT_SYNC_ENABLED is off"}

    opener = opener or login()
    today = timezone.localdate()
    date_from = today - timedelta(days=days_back)
    date_to = today + timedelta(days=days_ahead)
    tz = timezone.get_current_timezone()
    stats = {"workdays": 0, "events": 0, "entries": 0}
    shifts: dict[date, Shift] = {}

    for workday in parse_schedule_html(fetch_schedule_html(opener)):
        if not (date_from <= workday.date <= date_to):
            continue
        user = user_for_fio(workday.fio)
        shift = _shift_for_date(workday.date, workday.fio)
        shift.opened_by = shift.opened_by or user
        shift.opened_at = timezone.make_aware(datetime.combine(workday.date, workday.start), tz)
        if workday.date < today:
            shift.closed_by = shift.closed_by or user
            shift.closed_at = shift.closed_at or timezone.make_aware(
                datetime.combine(workday.date, workday.end), tz
            )
            shift.accepted = True
        else:
            # сегодня/будущее — плановые: приём/сдачу проставят события
            shift.closed_by = None
            shift.closed_at = None
        shift.save()
        shifts[workday.date] = shift
        stats["workdays"] += 1

    for event in parse_events_xlsx(fetch_events_xlsx(opener, date_from, date_to)):
        stats["events"] += 1
        day = timezone.localtime(event.occurred_at).date()
        shift = shifts.get(day) or _shift_for_date(day, event.fio)
        shifts[day] = shift
        if event.kind == "start":
            actual = user_for_fio(event.fio)
            if actual and shift.opened_by_id and shift.opened_by_id != actual.pk:
                # принял не тот, кто по графику (например, ведущий инженер подменил)
                if not shift.opening_notes:
                    shift.opening_notes = f"По графику: {shift.opened_by.full_name}"
            if actual:
                shift.opened_by = actual
            shift.opened_at = event.occurred_at
            shift.accepted = True
            shift.save(
                update_fields=["opened_by", "opened_at", "accepted", "opening_notes", "updated_at"]
            )
        elif event.kind == "end":
            shift.closed_by = user_for_fio(event.fio)
            shift.closed_at = event.occurred_at
            shift.accepted = True
            shift.save(
                update_fields=["closed_by", "closed_at", "accepted", "updated_at"]
            )
        else:
            data = parse_work_comment(event.comment)
            digest = hashlib.sha1(
                f"{event.occurred_at.isoformat()}|{event.fio}|{event.comment}".encode("utf-8")
            ).hexdigest()
            JournalEntry.objects.update_or_create(
                external_id=f"event:{digest}",
                defaults={
                    "shift": shift,
                    "entry_type": JournalEntry.EntryType.WORK,
                    "occurred_at": event.occurred_at,
                    "specialist": user_for_fio(event.fio),
                    "equipment_line": data["equipment_line"],
                    "downtime": data["downtime"],
                    "action_task": data["action_task"],
                    "solution": data["solution"],
                },
            )
            stats["entries"] += 1

    # Итоговый статус: будущее — «Планируется», прошлое — «Закрыта», сегодня — по факту
    now = timezone.now()
    today = timezone.localdate()
    for shift in shifts.values():
        if shift.date > today:
            new_status = Shift.Status.PLANNED
        elif shift.date < today:
            new_status = Shift.Status.CLOSED
        elif not shift.accepted:
            new_status = Shift.Status.PLANNED
        elif shift.closed_at and shift.closed_at <= now:
            new_status = Shift.Status.CLOSED
        else:
            new_status = Shift.Status.OPEN
        if shift.status != new_status:
            shift.status = new_status
            shift.save(update_fields=["status", "updated_at"])

    return stats


def quick_sync(force: bool = False) -> dict:
    """Лёгкий синк: один запрос к главной — обновляет статус текущей смены."""
    if not settings.SHIFT_SYNC_ENABLED and not force:
        return {"skipped": "SHIFT_SYNC_ENABLED is off"}
    if not force and cache.get("shift_quick_sync"):
        return {"skipped": "throttled"}
    cache.set("shift_quick_sync", 1, settings.SHIFT_SYNC_QUICK_TTL)

    try:
        opener = login()
        on_duty, fio = fetch_current_duty(opener)
    except Exception as exc:  # сеть/авторизация — не мешаем странице
        return {"error": str(exc)}

    if not on_duty:
        # дежурного нет — если смена была принята и открыта, значит сдали
        from shifts.models import Shift

        today = timezone.localdate()
        shift = Shift.objects.filter(
            external_id=f"workday:{today.isoformat()}",
            status=Shift.Status.OPEN,
            accepted=True,
        ).first()
        if shift:
            user = user_for_fio(fio)
            shift.status = Shift.Status.CLOSED
            shift.closed_at = timezone.now()
            if user:
                shift.closed_by = user
            shift.save(
                update_fields=["status", "closed_at", "closed_by", "updated_at"]
            )
            return {"on_duty": False, "closed": True}
        return {"on_duty": False}

    from shifts.models import Shift

    today = timezone.localdate()
    external_id = f"workday:{today.isoformat()}"
    shift = Shift.objects.filter(external_id=external_id).first()
    user = user_for_fio(fio)
    if shift is None:
        Shift.objects.create(
            external_id=external_id,
            date=today,
            kind=Shift.Kind.DAY,
            opened_by=user,
            opened_at=timezone.now(),
            status=Shift.Status.OPEN,
            accepted=True,
        )
    else:
        fields = []
        if user and shift.opened_by_id != user.pk:
            shift.opened_by = user
            fields.append("opened_by")
        if not shift.accepted:
            shift.accepted = True
            fields.append("accepted")
        if shift.status != Shift.Status.OPEN:
            shift.status = Shift.Status.OPEN
            fields.append("status")
        if fields:
            fields.append("updated_at")
            shift.save(update_fields=fields)
    return {"on_duty": True, "fio": fio}

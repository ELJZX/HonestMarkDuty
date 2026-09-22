from datetime import timedelta
from unittest import mock

from django.test import RequestFactory, TestCase
from django.urls import reverse
from django.utils import timezone

from accounts.models import User
from inventory.models import Condition, InventoryItem
from journal.models import JournalEntry, JournalExport
from shifts.forms import ShiftCheckForm
from shifts.models import Shift, ShiftCheck
from shifts.views import _plural_ru, close_shift_with, open_shift_for


class ShiftHelperTests(TestCase):
    def setUp(self):
        self.specialist = User.objects.create_user(
            username="spec", password="x", role=User.Role.SPECIALIST
        )

    def _request(self, user):
        request = RequestFactory().post("/")
        request.user = user
        return request

    def test_plural_ru_forms(self):
        self.assertEqual(_plural_ru(1, "запись", "записи", "записей"), "запись")
        self.assertEqual(_plural_ru(2, "запись", "записи", "записей"), "записи")
        self.assertEqual(_plural_ru(5, "запись", "записи", "записей"), "записей")
        self.assertEqual(_plural_ru(11, "запись", "записи", "записей"), "записей")
        self.assertEqual(_plural_ru(21, "запись", "записи", "записей"), "запись")

    def test_open_shift_for_creates_shift_and_marker(self):
        shift, message = open_shift_for(self._request(self.specialist))
        self.assertIsNotNone(shift)
        self.assertEqual(message, "Смена открыта")
        self.assertTrue(
            JournalEntry.objects.filter(
                shift=shift, entry_type=JournalEntry.EntryType.SHIFT_START
            ).exists()
        )

    def test_open_shift_for_twice_returns_none(self):
        Shift.objects.create(opened_by=self.specialist, status=Shift.Status.OPEN)
        shift, message = open_shift_for(self._request(self.specialist))
        self.assertIsNone(shift)
        self.assertIn("уже есть открытая смена", message)

    def test_close_shift_with_creates_end_marker(self):
        shift = Shift.objects.create(opened_by=self.specialist, status=Shift.Status.OPEN)
        close_shift_with(self._request(self.specialist), shift)
        shift.refresh_from_db()
        self.assertEqual(shift.status, Shift.Status.CLOSED)
        self.assertTrue(
            JournalEntry.objects.filter(
                shift=shift, entry_type=JournalEntry.EntryType.SHIFT_END
            ).exists()
        )

    def test_close_shift_with_tolerates_export_errors(self):
        shift = Shift.objects.create(opened_by=self.specialist, status=Shift.Status.OPEN)
        with mock.patch("shifts.views.export_full_journal", side_effect=RuntimeError("boom")), \
                mock.patch("checklists.services.finalize_shift_checklists",
                           side_effect=RuntimeError("boom")):
            close_shift_with(self._request(self.specialist), shift)
        shift.refresh_from_db()
        self.assertEqual(shift.status, Shift.Status.CLOSED)


class ShiftModelTests(TestCase):
    def test_str_and_is_open(self):
        shift = Shift.objects.create(status=Shift.Status.OPEN)
        self.assertIn("Смена", str(shift))
        self.assertTrue(shift.is_open)

    def test_duration_uses_closed_at(self):
        start = timezone.now() - timedelta(hours=8)
        shift = Shift.objects.create(
            status=Shift.Status.CLOSED, opened_at=start, closed_at=timezone.now()
        )
        self.assertGreaterEqual(shift.duration.total_seconds(), 8 * 3600 - 60)

    def test_duration_display_format(self):
        start = timezone.now()
        shift = Shift.objects.create(
            status=Shift.Status.CLOSED,
            opened_at=start,
            closed_at=start + timedelta(hours=2, minutes=3, seconds=4),
        )
        self.assertEqual(shift.duration_display, "02:03:04")

    def test_shiftcheck_str(self):
        shift = Shift.objects.create()
        item = InventoryItem.objects.create(name="Сканер")
        check = ShiftCheck.objects.create(shift=shift, item=item, wear_percent=10)
        self.assertIn("10", str(check))


class ShiftViewTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            username="admin", password="x", role=User.Role.ADMIN, is_superuser=True, is_staff=True
        )
        self.specialist = User.objects.create_user(
            username="spec", password="x", role=User.Role.SPECIALIST
        )
        self.item = InventoryItem.objects.create(name="Сканер", quantity=5)

    def test_list_and_open_page_render(self):
        self.client.force_login(self.specialist)
        self.assertEqual(self.client.get(reverse("shifts:shift_list")).status_code, 200)
        self.assertEqual(self.client.get(reverse("shifts:shift_open")).status_code, 200)

    def test_open_shift_creates_marker(self):
        self.client.force_login(self.specialist)
        response = self.client.post(
            reverse("shifts:shift_open"),
            {"date": timezone.localdate().isoformat(), "kind": Shift.Kind.DAY, "opening_notes": "старт"},
        )
        self.assertEqual(response.status_code, 302)
        shift = Shift.objects.open().first()
        self.assertEqual(shift.opened_by, self.specialist)
        self.assertTrue(
            JournalEntry.objects.filter(
                shift=shift, entry_type=JournalEntry.EntryType.SHIFT_START
            ).exists()
        )

    def test_open_shift_twice_redirects(self):
        self.client.force_login(self.specialist)
        Shift.objects.create(opened_by=self.specialist, status=Shift.Status.OPEN)
        response = self.client.post(
            reverse("shifts:shift_open"),
            {"date": timezone.localdate().isoformat(), "kind": Shift.Kind.DAY},
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Shift.objects.filter(opened_by=self.specialist).count(), 1)

    def test_close_shift_creates_end_marker_and_export(self):
        shift = Shift.objects.create(opened_by=self.specialist)
        JournalEntry.objects.create(shift=shift, action_task="Работа за смену")
        self.client.force_login(self.specialist)
        response = self.client.post(
            reverse("shifts:shift_close", args=[shift.pk]),
            {"equipment_condition": "норма", "inventory_notes": "ок", "closing_notes": "сдал", "handover_to": ""},
        )
        self.assertEqual(response.status_code, 302)
        shift.refresh_from_db()
        self.assertEqual(shift.status, Shift.Status.CLOSED)
        self.assertTrue(
            JournalEntry.objects.filter(
                shift=shift, entry_type=JournalEntry.EntryType.SHIFT_END
            ).exists()
        )
        self.assertTrue(JournalExport.objects.filter(is_full=True).exists())

    def test_close_shift_without_entries_still_exports_marker(self):
        shift = Shift.objects.create(opened_by=self.specialist)
        self.client.force_login(self.specialist)
        response = self.client.post(
            reverse("shifts:shift_close", args=[shift.pk]),
            {"equipment_condition": "", "inventory_notes": "", "closing_notes": "", "handover_to": ""},
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(JournalExport.objects.filter(is_full=True).exists())

    def test_add_check(self):
        shift = Shift.objects.create(opened_by=self.specialist)
        self.client.force_login(self.specialist)
        response = self.client.post(
            reverse("shifts:check_create", args=[shift.pk]),
            {"item": self.item.pk, "condition": Condition.GOOD, "wear_percent": 10, "comment": "ок"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(shift.checks.count(), 1)

    def test_add_check_to_closed_shift_rejected(self):
        shift = Shift.objects.create(status=Shift.Status.CLOSED, opened_by=self.specialist)
        self.client.force_login(self.specialist)
        self.client.post(
            reverse("shifts:check_create", args=[shift.pk]),
            {"item": self.item.pk, "condition": Condition.GOOD, "wear_percent": 10},
        )
        self.assertEqual(shift.checks.count(), 0)

    def test_check_form_valid(self):
        form = ShiftCheckForm(data={"item": self.item.pk, "condition": Condition.NEW, "wear_percent": 0})
        self.assertTrue(form.is_valid())

    def test_detail_and_close_pages_render(self):
        shift = Shift.objects.create(opened_by=self.specialist)
        JournalEntry.objects.create(shift=shift, action_task="x")
        self.client.force_login(self.specialist)
        detail = self.client.get(reverse("shifts:shift_detail", args=[shift.pk]))
        self.assertEqual(detail.status_code, 200)
        close = self.client.get(reverse("shifts:shift_close", args=[shift.pk]))
        self.assertEqual(close.status_code, 200)
        self.assertEqual(close.context["journal_count"], 1)

    def test_list_filters_by_status_and_date(self):
        closed = Shift.objects.create(
            opened_by=self.specialist, status=Shift.Status.CLOSED, date=timezone.localdate()
        )
        opened = Shift.objects.create(
            opened_by=self.specialist,
            status=Shift.Status.OPEN,
            date=timezone.localdate() - timedelta(days=1),
        )
        self.client.force_login(self.specialist)
        response = self.client.get(reverse("shifts:shift_list"), {"status": Shift.Status.CLOSED})
        self.assertIn(closed, response.context["shifts"])
        self.assertNotIn(opened, response.context["shifts"])
        response = self.client.get(reverse("shifts:shift_list"), {"date": closed.date.isoformat()})
        self.assertIn(closed, response.context["shifts"])
        self.assertNotIn(opened, response.context["shifts"])

    def test_add_check_invalid_form_shows_error(self):
        shift = Shift.objects.create(opened_by=self.specialist)
        self.client.force_login(self.specialist)
        response = self.client.post(reverse("shifts:check_create", args=[shift.pk]), {})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(shift.checks.count(), 0)

    def test_shift_open_ajax_requires_login(self):
        response = self.client.post(reverse("shifts:shift_open_ajax"))
        self.assertEqual(response.status_code, 403)

    def test_shift_open_ajax_forbidden_for_viewer(self):
        viewer = User.objects.create_user(username="viewer", password="x", role=User.Role.VIEWER)
        self.client.force_login(viewer)
        response = self.client.post(reverse("shifts:shift_open_ajax"))
        self.assertEqual(response.status_code, 403)

    def test_shift_open_ajax_opens_shift(self):
        self.client.force_login(self.specialist)
        response = self.client.post(reverse("shifts:shift_open_ajax"))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["ok"])
        self.assertTrue(Shift.objects.open().filter(opened_by=self.specialist).exists())

    def test_shift_open_ajax_duplicate_returns_400(self):
        Shift.objects.create(opened_by=self.specialist, status=Shift.Status.OPEN)
        self.client.force_login(self.specialist)
        response = self.client.post(reverse("shifts:shift_open_ajax"))
        self.assertEqual(response.status_code, 400)
        self.assertFalse(response.json()["ok"])

    def test_shift_close_ajax_requires_login(self):
        shift = Shift.objects.create(opened_by=self.specialist)
        response = self.client.post(reverse("shifts:shift_close_ajax", args=[shift.pk]))
        self.assertEqual(response.status_code, 403)

    def test_shift_close_ajax_forbidden_for_viewer(self):
        shift = Shift.objects.create(opened_by=self.specialist)
        viewer = User.objects.create_user(username="viewer2", password="x", role=User.Role.VIEWER)
        self.client.force_login(viewer)
        response = self.client.post(reverse("shifts:shift_close_ajax", args=[shift.pk]))
        self.assertEqual(response.status_code, 403)

    def test_shift_close_ajax_closes_shift(self):
        shift = Shift.objects.create(opened_by=self.specialist, status=Shift.Status.OPEN)
        self.client.force_login(self.specialist)
        response = self.client.post(reverse("shifts:shift_close_ajax", args=[shift.pk]))
        self.assertEqual(response.status_code, 200)
        shift.refresh_from_db()
        self.assertEqual(shift.status, Shift.Status.CLOSED)


class ShiftOpenEdgeTests(TestCase):
    def test_open_form_valid_redirects_to_list_when_shift_exists(self):
        specialist = User.objects.create_user(
            username="shift-open-edge", password="x", role=User.Role.SPECIALIST
        )
        self.client.force_login(specialist)
        with mock.patch(
            "shifts.views.open_shift_for", return_value=(None, "У вас уже есть открытая смена.")
        ):
            response = self.client.post(reverse("shifts:shift_open"))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("shifts:shift_list"))


class ShiftIntegrationParseTests(TestCase):
    def test_parse_schedule_html(self):
        from shifts.integration import parse_schedule_html

        html = (
            "<table><tbody>"
            "<tr><th>День</th><th>ФИО</th><th>Начало</th><th>Конец</th></tr>"
            "<tr><td>18 сентября 2026 г.</td><td>Бородин Александр Владимирович</td>"
            "<td>8:00</td><td>20:00</td><td>+7</td><td></td></tr>"
            "</tbody></table>"
        )
        workdays = parse_schedule_html(html)
        self.assertEqual(len(workdays), 1)
        self.assertEqual(workdays[0].date.isoformat(), "2026-09-18")
        self.assertEqual(workdays[0].fio, "Бородин Александр Владимирович")
        self.assertEqual(workdays[0].start.hour, 8)
        self.assertEqual(workdays[0].end.hour, 20)

    def test_parse_work_comment(self):
        from shifts.integration import parse_work_comment

        data = parse_work_comment("16:49\tFinnah\t\tнастроил датчик\tпроверил")
        self.assertEqual(data["equipment_line"], "Finnah")
        self.assertEqual(data["action_task"], "настроил датчик")
        self.assertEqual(data["solution"], "проверил")
        self.assertEqual(data["downtime"], "16:49")

    def test_parse_events_xlsx(self):
        import io

        from openpyxl import Workbook

        from shifts.integration import parse_events_xlsx

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Дата и время", "ФИО", "Тип события", "Комментарий"])
        sheet.append(["18.09.2026 08:21:13", "Бородин Александр Владимирович", "Начало смены", "Смена принята"])
        sheet.append(["18.09.2026 19:59:40", "Бородин Александр Владимирович", "Запись", "16:49\tFinnah\t\tнастроил\tок"])
        sheet.append(["18.09.2026 20:02:00", "Бородин Александр Владимирович", "Конец смены", "Смена сдана"])
        buffer = io.BytesIO()
        workbook.save(buffer)

        events = parse_events_xlsx(buffer.getvalue())
        self.assertEqual(len(events), 3)
        self.assertEqual(events[0].kind, "start")
        self.assertEqual(events[1].kind, "work")
        self.assertEqual(events[2].kind, "end")


class ShiftIntegrationSyncTests(TestCase):
    def _xlsx(self, day):
        import io

        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Дата и время", "ФИО", "Тип события", "Комментарий"])
        sheet.append([f"{day:%d.%m.%Y} 08:21:13", "Бородин Александр Владимирович", "Начало смены", "принят"])
        sheet.append([f"{day:%d.%m.%Y} 19:59:40", "Бородин Александр Владимирович", "Запись", "16:49\tFinnah\t\tнастроил\tок"])
        sheet.append([f"{day:%d.%m.%Y} 20:02:00", "Бородин Александр Владимирович", "Конец смены", "сдан"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def _schedule_html(self, day):
        from shifts.integration import MONTHS

        month_name = {v: k for k, v in MONTHS.items()}[day.month]
        return (
            "<table><tr><td>"
            f"{day.day} {month_name} {day.year} г."
            "</td><td>Бородин Александр Владимирович</td><td>8:00</td><td>20:00</td>"
            "<td>+7</td><td></td></tr></table>"
        )

    def test_sync_creates_shift_and_journal(self):
        from shifts.integration import sync_window

        day = timezone.localdate() - timedelta(days=2)
        with mock.patch("shifts.integration.fetch_schedule_html", return_value=self._schedule_html(day)), \
                mock.patch("shifts.integration.fetch_events_xlsx", return_value=self._xlsx(day)):
            result = sync_window(30, 7, opener=object(), force=True)

        shift = Shift.objects.get(external_id=f"workday:{day.isoformat()}")
        self.assertEqual(shift.opened_by.last_name, "Бородин")
        self.assertEqual(shift.status, Shift.Status.CLOSED)
        self.assertIsNotNone(shift.closed_at)
        self.assertEqual(result["entries"], 1)
        self.assertTrue(JournalEntry.objects.filter(equipment_line="Finnah").exists())

    def test_sync_disabled_is_skipped(self):
        from shifts.integration import sync_window

        result = sync_window(7, 7, opener=object())
        self.assertIn("skipped", result)

    def test_quick_sync_marks_today_open(self):
        from django.core.cache import cache

        from shifts.integration import quick_sync

        cache.clear()
        with mock.patch("shifts.integration.login", return_value=object()), mock.patch(
            "shifts.integration.fetch_current_duty",
            return_value=(True, "Бородин Александр Владимирович"),
        ):
            result = quick_sync(force=True)
        self.assertTrue(result["on_duty"])
        shift = Shift.objects.get(
            external_id=f"workday:{timezone.localdate().isoformat()}"
        )
        self.assertEqual(shift.status, Shift.Status.OPEN)
        self.assertEqual(shift.opened_by.last_name, "Бородин")

    def test_quick_sync_throttled(self):
        from django.core.cache import cache

        from shifts.integration import quick_sync

        cache.clear()
        cache.set("shift_quick_sync", 1, 120)
        result = quick_sync()
        self.assertIn("skipped", result)

    def test_shift_list_triggers_quick_sync(self):
        from django.core.cache import cache

        cache.clear()
        specialist = User.objects.create_user(
            username="quick-sync-user", password="x", role=User.Role.SPECIALIST
        )
        self.client.force_login(specialist)
        with mock.patch("shifts.integration.quick_sync") as patched:
            self.client.get(reverse("shifts:shift_list"))
        patched.assert_called_once()

    def test_sync_reflects_substitute(self):
        from shifts.integration import MONTHS, sync_window

        day = timezone.localdate() - timedelta(days=2)
        month_name = {v: k for k, v in MONTHS.items()}[day.month]
        schedule = (
            "<table><tr><td>"
            f"{day.day} {month_name} {day.year} г."
            "</td><td>Иванов Иван Иванович</td><td>8:00</td><td>20:00</td></tr></table>"
        )
        import io

        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Дата и время", "ФИО", "Тип события", "Комментарий"])
        sheet.append([f"{day:%d.%m.%Y} 08:05:00", "Петров Пётр Петрович", "Начало смены", "принят"])
        sheet.append([f"{day:%d.%m.%Y} 20:01:00", "Петров Пётр Петрович", "Конец смены", "сдан"])
        buffer = io.BytesIO()
        workbook.save(buffer)

        with mock.patch("shifts.integration.fetch_schedule_html", return_value=schedule), \
                mock.patch("shifts.integration.fetch_events_xlsx", return_value=buffer.getvalue()):
            sync_window(30, 7, opener=object(), force=True)

        shift = Shift.objects.get(external_id=f"workday:{day.isoformat()}")
        self.assertEqual(shift.opened_by.last_name, "Петров")
        self.assertIn("Иванов", shift.opening_notes)
        self.assertEqual(shift.closed_by.last_name, "Петров")
